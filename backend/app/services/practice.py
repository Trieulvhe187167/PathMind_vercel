"""Question bank helpers: generating questions with AI, grading answers, spaced repetition (SM-2)."""

from __future__ import annotations

import csv
from datetime import timedelta
import io
import json
import re
from typing import Any

from sqlalchemy.orm import Session

from ..models import Question, ReviewState, User, utcnow
from . import llm

TYPES = ("mcq", "multi", "true_false", "short", "essay")
LANG = {"vi": "Vietnamese", "en": "English", "zh": "Simplified Chinese"}


def question_public(q: Question, with_answer: bool = True) -> dict:
    out = {"id": q.id, "course_id": q.course_id, "topic": q.topic, "type": q.type, "prompt": q.prompt,
           "options": q.options or [], "difficulty": q.difficulty, "tags": q.tags or [], "source": q.source or {},
           "created_at": q.created_at.isoformat() + "Z" if q.created_at else None}
    if with_answer:
        out["answer"] = q.answer or {}
        out["explanation"] = q.explanation
    return out


def clean(item: Any) -> dict | None:
    """Validate one question as produced by the AI or an import; None if unusable."""
    if not isinstance(item, dict):
        return None
    qtype = str(item.get("type") or "mcq").lower().replace("-", "_")
    qtype = {"single": "mcq", "multiple_choice": "mcq", "tf": "true_false", "truefalse": "true_false"}.get(qtype, qtype)
    if qtype not in TYPES:
        return None
    prompt = str(item.get("prompt") or item.get("question") or "").strip()
    if not prompt:
        return None
    options = [str(o).strip() for o in (item.get("options") or []) if str(o).strip()]
    ans = item.get("answer")
    answer: dict = {}
    if qtype == "true_false":
        options = options or ["True", "False"]
        val = ans if not isinstance(ans, dict) else ans.get("correct")
        if isinstance(val, list):
            val = val[0] if val else None
        if isinstance(val, str):
            val = 0 if val.strip().lower() in ("true", "đúng", "t", "0", "a", "对") else 1
        answer = {"correct": [int(val) if val in (0, 1) else 0]}
    elif qtype in ("mcq", "multi"):
        if len(options) < 2:
            return None
        raw = ans.get("correct") if isinstance(ans, dict) else ans
        idx: list[int] = []
        for v in raw if isinstance(raw, list) else [raw]:
            if isinstance(v, int) and 0 <= v < len(options):
                idx.append(v)
            elif isinstance(v, str):
                s = v.strip()
                if len(s) == 1 and s.upper() in "ABCDEFGH" and "ABCDEFGH".index(s.upper()) < len(options):
                    idx.append("ABCDEFGH".index(s.upper()))
                elif s.isdigit() and int(s) < len(options):
                    idx.append(int(s))
                elif s in options:
                    idx.append(options.index(s))
        if not idx:
            return None
        answer = {"correct": sorted(set(idx)) if qtype == "multi" else [idx[0]]}
    else:
        text = ans.get("text") if isinstance(ans, dict) else ans
        answer = {"text": str(text or "").strip()}
    try:
        diff = max(1, min(5, int(item.get("difficulty") or 2)))
    except (TypeError, ValueError):
        diff = 2
    return {"type": qtype, "prompt": prompt[:5000], "options": options[:8], "answer": answer,
            "explanation": str(item.get("explanation") or "")[:5000], "difficulty": diff,
            "topic": str(item.get("topic") or "")[:200]}


def generate(db: Session, user: User, *, topic: str, count: int, qtype: str = "mixed", difficulty: int = 2,
             material: str = "") -> list[dict]:
    kinds = "multiple choice (mcq), true/false and short answer" if qtype == "mixed" else qtype
    prompt = (f"Write {count} practice questions about: {topic}. Types: {kinds}. Difficulty {difficulty}/5. "
              f"Write them in {LANG.get(user.locale, 'Vietnamese')}. "
              "Respond with JSON: {\"questions\": [{\"type\": \"mcq|multi|true_false|short\", \"prompt\": \"...\", "
              "\"options\": [\"...\"], \"answer\": {\"correct\": [0]} or {\"text\": \"...\"}, \"explanation\": \"...\", "
              "\"difficulty\": 1-5}]}")
    if material:
        prompt += "\n\nBase the questions on this material:\n" + material[:12000]
    data = llm.complete_json(db, user=user, messages=[{"role": "user", "content": prompt}], purpose="quiz_generate",
                             fallback={})
    items = data.get("questions") if isinstance(data, dict) else data if isinstance(data, list) else []
    out = []
    for it in items or []:
        c = clean(it)
        if c:
            c["topic"] = c["topic"] or topic[:200]
            out.append(c)
    return out[:count]


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[^\w\s]", "", s.lower())).strip()


def grade(db: Session, user: User, q: Question, response: dict) -> tuple[bool | None, float | None, str]:
    """(correct, score 0..1, feedback). AI grades free text when an answer key alone is not enough."""
    ans = q.answer or {}
    if q.type in ("mcq", "true_false", "multi"):
        chosen = response.get("choices")
        if not isinstance(chosen, list):
            chosen = [response.get("choice")]
        chosen = sorted({int(c) for c in chosen if isinstance(c, int) or (isinstance(c, str) and c.isdigit())})
        correct = sorted(ans.get("correct") or [])
        ok = chosen == correct
        if q.type == "multi" and not ok and correct:
            hit = len(set(chosen) & set(correct))
            wrong = len(set(chosen) - set(correct))
            score = max(0.0, (hit - wrong) / len(correct))
            return False, round(score, 2), ""
        return ok, 1.0 if ok else 0.0, ""
    text = str(response.get("text") or "").strip()
    key = str(ans.get("text") or "")
    if not text:
        return False, 0.0, ""
    if q.type == "short" and key and _norm(text) == _norm(key):
        return True, 1.0, ""
    prompt = (f"Grade the learner's answer. Question: {q.prompt}\nReference answer: {key or '(none)'}\n"
              f"Learner's answer: {text}\nRespond with JSON: {{\"score\": 0.0-1.0, \"feedback\": \"one or two sentences "
              f"in {LANG.get(user.locale, 'Vietnamese')}\"}}")
    data = llm.complete_json(db, user=user, messages=[{"role": "user", "content": prompt}], purpose="grade", fallback={})
    try:
        score = max(0.0, min(1.0, float((data or {}).get("score"))))
    except (TypeError, ValueError):
        # No usable AI grade: fall back to word overlap with the reference answer.
        a, b = set(_norm(text).split()), set(_norm(key).split())
        score = round(len(a & b) / len(b), 2) if b else 0.0
        data = {}
    return score >= 0.6, round(score, 2), str((data or {}).get("feedback") or "")


def schedule(state: ReviewState, quality: int) -> None:
    """SM-2: quality 0..5. Below 3 = forgotten (start over)."""
    now = utcnow()
    if quality < 3:
        state.repetitions = 0
        state.lapses += 1
        state.interval_days = 0.0
        due = now + timedelta(minutes=10)  # see it again in this session
    else:
        state.repetitions += 1
        if state.repetitions == 1:
            state.interval_days = 1
        elif state.repetitions == 2:
            state.interval_days = 6
        else:
            state.interval_days = round(state.interval_days * state.ease, 1)
        due = now + timedelta(days=state.interval_days)
    state.ease = max(1.3, state.ease + (0.1 - (5 - quality) * (0.08 + (5 - quality) * 0.02)))
    state.due_at, state.last_reviewed_at = due, now


def parse_import(filename: str, data: bytes) -> list[dict]:
    name = filename.lower()
    if name.endswith(".json"):
        js = json.loads(data.decode("utf-8-sig"))
        items = js.get("questions") if isinstance(js, dict) else js
        return [c for c in (clean(i) for i in items or []) if c]
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        rows = [["" if v is None else str(v).strip() for v in r] for r in ws.iter_rows(values_only=True)]
    elif name.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), dialect)]
    else:
        raise ValueError("Use a .csv, .xlsx or .json file.")
    rows = [r for r in rows if any(r)]
    if not rows:
        return []
    head = [h.lower() for h in rows[0]]
    out = []
    for r in rows[1:]:
        d = dict(zip(head, r))
        options = [d[k] for k in ("option_a", "option_b", "option_c", "option_d", "option_e") if d.get(k)]
        if not options and d.get("options"):
            options = [o.strip() for o in d["options"].split("|") if o.strip()]
        answer: Any = d.get("answer", "")
        if (d.get("type") or "mcq") in ("multi",) and "," in answer:
            answer = {"correct": [a.strip() for a in answer.split(",")]}
        c = clean({"type": d.get("type") or ("mcq" if options else "short"), "prompt": d.get("prompt") or d.get("question"),
                   "options": options, "answer": answer, "explanation": d.get("explanation"), "topic": d.get("topic"),
                   "difficulty": d.get("difficulty")})
        if c:
            out.append(c)
    return out
