"""Bulk import of learners + plans from CSV / Excel (preview first, then apply).

One row per person: email (required), full name, username, role (learner/educator), plan,
days or expiry date, group. Existing e-mails get the plan and group; new ones get an account
with a temporary password or an invitation link, as the admin chooses.
"""

from __future__ import annotations

import csv
from datetime import date, datetime
import io
import re
from typing import Any

from email_validator import EmailNotValidError, validate_email
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..models import BulkImport, BulkImportRow, LearnerGroup, LearnerGroupMember, Plan, User, UserInvitation, utcnow
from ..security import hash_password
from . import accounts, billing, invitations

COLUMNS = ["email", "full_name", "username", "role", "plan", "days", "expires_on", "group"]
SYNONYMS = {
    "e-mail": "email", "mail": "email", "địa chỉ email": "email",
    "name": "full_name", "full name": "full_name", "họ tên": "full_name", "ho ten": "full_name", "họ và tên": "full_name", "姓名": "full_name",
    "user": "username", "tên đăng nhập": "username", "用户名": "username",
    "vai trò": "role", "角色": "role",
    "gói": "plan", "goi": "plan", "plan id": "plan", "套餐": "plan",
    "số ngày": "days", "so ngay": "days", "天数": "days",
    "ngày hết hạn": "expires_on", "expiry": "expires_on", "expires": "expires_on", "end date": "expires_on", "到期日": "expires_on",
    "nhóm": "group", "nhom": "group", "group name": "group", "分组": "group",
}
MAX_ROWS = 5000


class ImportError_(ValueError):
    pass


def _norm_header(h: Any) -> str:
    key = re.sub(r"\s+", " ", str(h or "").strip().lower())
    return SYNONYMS.get(key, key.replace(" ", "_"))


def read_table(filename: str, data: bytes) -> list[dict[str, str]]:
    name = filename.lower()
    rows: list[list[Any]]
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise ImportError_("Could not open the Excel file.") from exc
        ws = wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    elif name.endswith((".csv", ".txt")):
        text = data.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
    else:
        raise ImportError_("Upload a .csv or .xlsx file.")
    rows = [r for r in rows if any(str(c or "").strip() for c in r)]
    if not rows:
        raise ImportError_("The file is empty.")
    header = [_norm_header(h) for h in rows[0]]
    if "email" not in header:
        raise ImportError_("The first row must be the column names, including 'email'.")
    out = []
    for r in rows[1:MAX_ROWS + 1]:
        item = {}
        for i, h in enumerate(header):
            if h and i < len(r):
                v = r[i]
                if isinstance(v, (datetime, date)):
                    v = v.strftime("%Y-%m-%d")
                elif isinstance(v, float) and v.is_integer():
                    v = int(v)
                item[h] = "" if v is None else str(v).strip()
        out.append(item)
    return out


def _parse_date(s: str) -> date | None:
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def preview(db: Session, imp: BulkImport, rows: list[dict[str, str]]) -> None:
    """Validate every row and store the planned action. Caller commits."""
    opts = imp.options or {}
    default_plan, default_days = opts.get("default_plan") or "", opts.get("default_days")
    plans = {p.id for p in db.scalars(select(Plan))}
    seen: set[str] = set()
    ok = err = 0
    today = utcnow().date()
    for n, data in enumerate(rows, start=2):  # row 1 is the header
        errors: list[str] = []
        raw_email = data.get("email", "")
        try:
            email = validate_email(raw_email, check_deliverability=False).normalized.lower()
        except EmailNotValidError:
            email = ""
            errors.append("invalid e-mail")
        if email in seen:
            errors.append("e-mail appears twice in the file")
        seen.add(email)
        role = (data.get("role") or "learner").lower()
        if role not in ("learner", "educator"):
            errors.append("role must be learner or educator")
        plan = (data.get("plan") or default_plan or "").strip().lower()
        if plan and plan not in plans:
            errors.append(f"unknown plan '{plan}'")
        days: int | None = None
        if data.get("expires_on"):
            d = _parse_date(data["expires_on"])
            if d is None:
                errors.append("expiry date must be YYYY-MM-DD or DD/MM/YYYY")
            elif d <= today:
                errors.append("expiry date is in the past")
            else:
                days = (d - today).days
        elif data.get("days"):
            try:
                days = int(float(data["days"]))
                if days < 0:
                    raise ValueError
            except ValueError:
                errors.append("days must be a whole number (0 = never expires)")
        elif plan:
            days = int(default_days) if default_days not in (None, "") else 30
        username = data.get("username", "")
        if username and not accounts.USERNAME_RE.match(username):
            errors.append("username: 3-32 letters, numbers, dot, dash or underscore")
        existing = db.scalar(select(User).where(User.email == email)) if email else None
        action = "create_user"
        status = "pending"
        if existing is not None:
            if existing.role in billing.UNLIMITED_ROLES:
                action, status = "skip", "skipped"
                errors.append("R&D / admin account: not changed")
            else:
                action = "update_user" if (plan or data.get("group")) else "skip"
                if action == "skip":
                    status = "skipped"
        elif username and db.scalar(select(User.id).where(func.lower(User.username) == username.lower())):
            errors.append("username is taken")
        if errors and status != "skipped":
            status = "error"
        row = BulkImportRow(import_id=imp.id, row_number=n,
                            data={**data, "email": email or raw_email, "role": role, "plan": plan, "days": days},
                            action=action, status=status, error="; ".join(errors)[:500],
                            user_id=existing.id if existing else None)
        db.add(row)
        if status == "error":
            err += 1
        elif status == "pending":
            ok += 1
    imp.total_rows, imp.ok_rows, imp.error_rows = len(rows), ok, err
    imp.status = "previewed"


def _unique_username(db: Session, base: str) -> str:
    base = re.sub(r"[^a-z0-9._-]", "", base.lower())[:24] or "user"
    if len(base) < 3:
        base = base + "001"
    name, n = base, 1
    while db.scalar(select(User.id).where(func.lower(User.username) == name)):
        n += 1
        name = f"{base}{n}"
    return name


def _group(db: Session, name: str, admin_id: str, cache: dict[str, LearnerGroup]) -> LearnerGroup:
    key = name.strip().lower()
    if key in cache:
        return cache[key]
    g = db.scalar(select(LearnerGroup).where(func.lower(LearnerGroup.name) == key))
    if g is None:
        g = LearnerGroup(name=name.strip()[:120], created_by=admin_id)
        db.add(g)
        db.flush()
    cache[key] = g
    return g


def apply(db: Session, imp: BulkImport, admin: User) -> dict:
    """Carry out every pending row. Returns one-time credentials (passwords or invitation tokens)."""
    opts = imp.options or {}
    mode = opts.get("mode", "password")
    default_group = db.get(LearnerGroup, opts["group_id"]) if opts.get("group_id") else None
    groups: dict[str, LearnerGroup] = {}
    created = updated = plans_given = invites = 0
    credentials: list[dict] = []
    imp.status = "applying"
    rows = db.scalars(select(BulkImportRow).where(BulkImportRow.import_id == imp.id, BulkImportRow.status == "pending")
                      .order_by(BulkImportRow.row_number)).all()
    for row in rows:
        d = row.data or {}
        try:
            with db.begin_nested():
                user = db.get(User, row.user_id) if row.user_id else None
                if user is None and row.action == "create_user":
                    if db.scalar(select(User.id).where(User.email == d["email"])):
                        raise ValueError("e-mail was registered after the preview")
                    if mode == "invite":
                        token, token_hash = invitations.new_token()
                        if invitations.open_for_email(db, d["email"]):
                            raise ValueError("an invitation to this e-mail is still open")
                        inv = UserInvitation(email=d["email"], role=d.get("role", "learner"), token_hash=token_hash,
                                             invited_by=admin.id, expires_at=invitations.expiry(int(opts.get("invite_days") or 14)))
                        db.add(inv)
                        db.flush()
                        credentials.append({"row": row.row_number, "email": d["email"], "token": token})
                        invites += 1
                        # The plan and group are applied when the invitation is accepted.
                        inv_plan = {"plan": d.get("plan"), "days": d.get("days"), "group": d.get("group") or (default_group.name if default_group else "")}
                        row.data = {**d, "invitation_id": inv.id, "pending_on_accept": inv_plan}
                        row.action, row.status = "invite", "ok"
                        continue
                    pw = accounts.temporary_password()
                    user = User(username=_unique_username(db, d.get("username") or d["email"].split("@")[0]),
                                email=d["email"], full_name=(d.get("full_name") or "")[:120], role=d.get("role", "learner"),
                                password_hash=hash_password(pw), must_change_password=True, created_by=admin.id)
                    db.add(user)
                    db.flush()
                    credentials.append({"row": row.row_number, "email": user.email, "username": user.username, "password": pw})
                    row.user_id = user.id
                    created += 1
                elif user is not None:
                    updated += 1
                if user is None:
                    raise ValueError("account not found")
                if d.get("plan"):
                    billing.activate(db, user.id, d["plan"], days=d.get("days") if d.get("days") is not None else 30,
                                     source="import", bulk_import_id=imp.id)
                    plans_given += 1
                gname = d.get("group") or (default_group.name if default_group else "")
                if gname:
                    g = default_group if (default_group and gname == default_group.name) else _group(db, gname, admin.id, groups)
                    if db.get(LearnerGroupMember, (g.id, user.id)) is None:
                        db.add(LearnerGroupMember(group_id=g.id, user_id=user.id, added_by_import_id=imp.id))
                row.status = "ok"
        except Exception as exc:  # noqa: BLE001 - recorded on the row
            row.status, row.error = "error", str(exc)[:500]
    imp.ok_rows = db.scalar(select(func.count()).select_from(BulkImportRow).where(
        BulkImportRow.import_id == imp.id, BulkImportRow.status == "ok")) or 0
    imp.error_rows = db.scalar(select(func.count()).select_from(BulkImportRow).where(
        BulkImportRow.import_id == imp.id, BulkImportRow.status == "error")) or 0
    imp.summary = {"created": created, "updated": updated, "plans_assigned": plans_given, "invitations": invites}
    imp.status, imp.applied_at = "applied", utcnow()
    return {"summary": imp.summary, "credentials": credentials, "mode": mode}


def apply_pending_invitation(db: Session, inv: UserInvitation, user: User) -> None:
    """When an imported invitation is accepted, give the plan and group it carried."""
    row = db.scalar(select(BulkImportRow).where(BulkImportRow.action == "invite",
                                                BulkImportRow.user_id.is_(None)).where(
        BulkImportRow.data["invitation_id"].as_string() == inv.id))
    if row is None:
        return
    pending = (row.data or {}).get("pending_on_accept") or {}
    if pending.get("plan"):
        billing.activate(db, user.id, pending["plan"], days=pending.get("days") if pending.get("days") is not None else 30,
                         source="import", bulk_import_id=row.import_id)
    if pending.get("group"):
        g = db.scalar(select(LearnerGroup).where(func.lower(LearnerGroup.name) == pending["group"].lower()))
        if g is None:
            g = LearnerGroup(name=pending["group"][:120], created_by=inv.invited_by)
            db.add(g)
            db.flush()
        if db.get(LearnerGroupMember, (g.id, user.id)) is None:
            db.add(LearnerGroupMember(group_id=g.id, user_id=user.id, added_by_import_id=row.import_id))
    row.user_id = user.id


def template_csv() -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(COLUMNS)
    w.writerow(["an.nguyen@company.com", "Nguyễn Văn An", "", "learner", "plus", "90", "", "Sales team"])
    w.writerow(["binh.tran@company.com", "Trần Bình", "binhtran", "educator", "pro", "", "2027-06-30", "Trainers"])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def template_xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Learners"
    ws.append(COLUMNS)
    ws.append(["an.nguyen@company.com", "Nguyễn Văn An", "", "learner", "plus", 90, "", "Sales team"])
    ws.append(["binh.tran@company.com", "Trần Bình", "binhtran", "educator", "pro", "", "2027-06-30", "Trainers"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def rows_csv(rows: list[BulkImportRow], only_errors: bool = False) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["row", "status", "action", "error"] + COLUMNS)
    for r in rows:
        if only_errors and r.status != "error":
            continue
        d = r.data or {}
        w.writerow([r.row_number, r.status, r.action, r.error] + [d.get(c, "") if d.get(c) is not None else "" for c in COLUMNS])
    return ("﻿" + buf.getvalue()).encode("utf-8")
