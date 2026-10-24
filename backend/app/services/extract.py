"""Turn uploaded documents into plain text with headings kept as markdown ``#`` lines.

Supported: .txt .md .csv .json .html, .docx (standard library), .pdf (needs ``pypdf``),
.xlsx (``openpyxl``). Returns pages so citations can point at a page when the format has them.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from html.parser import HTMLParser
import io
import re
import zipfile
from xml.etree import ElementTree


class ExtractError(ValueError):
    pass


@dataclass
class Extracted:
    text: str
    pages: list[str] = field(default_factory=list)  # one string per page when known

    @property
    def page_count(self) -> int:
        return len(self.pages)


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1258", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


class _HTMLText(HTMLParser):
    BLOCK = {"p", "div", "br", "li", "tr", "section", "article", "table", "ul", "ol"}

    def __init__(self) -> None:
        super().__init__()
        self.out: list[str] = []
        self.skip = 0
        self.heading: int | None = None

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "noscript"):
            self.skip += 1
        elif re.fullmatch(r"h[1-6]", tag):
            self.heading = int(tag[1])
            self.out.append("\n\n" + "#" * self.heading + " ")
        elif tag in self.BLOCK:
            self.out.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style", "noscript"):
            self.skip = max(0, self.skip - 1)
        elif re.fullmatch(r"h[1-6]", tag):
            self.heading = None
            self.out.append("\n")
        elif tag in self.BLOCK:
            self.out.append("\n")

    def handle_data(self, data):
        if not self.skip:
            self.out.append(data)


def html_to_text(html: str) -> str:
    p = _HTMLText()
    p.feed(html)
    text = "".join(p.out)
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def docx_to_text(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            xml = z.read("word/document.xml")
    except (zipfile.BadZipFile, KeyError) as exc:
        raise ExtractError("Not a valid .docx file.") from exc
    root = ElementTree.fromstring(xml)
    lines: list[str] = []
    for para in root.iter(f"{_W}p"):
        style = para.find(f"{_W}pPr/{_W}pStyle")
        level = 0
        if style is not None:
            m = re.search(r"(?:Heading|heading|Title)(\d?)", style.get(f"{_W}val", ""))
            if m:
                level = int(m.group(1) or 1)
        text = "".join(t.text or "" for t in para.iter(f"{_W}t")).strip()
        if not text:
            continue
        lines.append(("#" * min(level, 6) + " " + text) if level else text)
    return "\n\n".join(lines)


def pdf_pages(data: bytes) -> list[str]:
    try:
        from pypdf import PdfReader  # optional dependency
    except ImportError as exc:  # pragma: no cover - depends on the install
        raise ExtractError("PDF support needs the 'pypdf' package (run setup.bat).") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
        return [(page.extract_text() or "").strip() for page in reader.pages]
    except Exception as exc:  # noqa: BLE001 - pypdf raises many types for broken files
        raise ExtractError(f"Could not read the PDF: {exc}") from exc


def xlsx_to_text(data: bytes) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ExtractError("Excel support needs the 'openpyxl' package.") from exc
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"# {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract(filename: str, data: bytes, mime_type: str = "") -> Extracted:
    name = (filename or "").lower()
    if name.endswith(".pdf") or mime_type == "application/pdf":
        pages = pdf_pages(data)
        return Extracted("\n\n".join(pages), pages)
    if name.endswith(".docx"):
        return Extracted(docx_to_text(data))
    if name.endswith(".xlsx"):
        return Extracted(xlsx_to_text(data))
    if name.endswith((".html", ".htm")) or mime_type.startswith("text/html"):
        return Extracted(html_to_text(_decode(data)))
    if name.endswith((".txt", ".md", ".markdown", ".csv", ".json", ".rst", ".tex")) or mime_type.startswith("text/"):
        return Extracted(_decode(data))
    raise ExtractError("Unsupported file type. Use PDF, Word (.docx), Excel (.xlsx), HTML, Markdown or text.")


# ── Splitting text into passages for retrieval ─────────────────────────────


@dataclass
class Passage:
    seq: int
    heading_path: str
    text: str
    page_from: int | None = None
    page_to: int | None = None


_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")


def split_passages(ex: Extracted, max_chars: int = 1200, overlap: int = 150) -> list[Passage]:
    """Cut text into ~max_chars passages that never cross a heading, remembering the heading path."""
    units: list[tuple[int | None, str]] = []
    if ex.pages:
        for i, page in enumerate(ex.pages, start=1):
            units.append((i, page))
    else:
        units.append((None, ex.text))

    passages: list[Passage] = []
    path: list[str] = []
    buf: list[str] = []
    buf_pages: list[int] = []

    def flush() -> None:
        text = "\n".join(buf).strip()
        if text:
            pages = [p for p in buf_pages if p is not None]
            passages.append(Passage(len(passages) + 1, " > ".join(path), text,
                                    min(pages) if pages else None, max(pages) if pages else None))
        buf.clear()
        buf_pages.clear()

    def pieces(line: str) -> list[str]:
        """Break a very long line (a paragraph without newlines) at spaces."""
        if len(line) <= max_chars:
            return [line]
        out, cur = [], ""
        for word in line.split(" "):
            if cur and len(cur) + len(word) + 1 > max_chars:
                out.append(cur)
                cur = word
            else:
                cur = f"{cur} {word}" if cur else word
        if cur:
            out.append(cur)
        return out

    for page, text in units:
        for line in (p for raw in text.splitlines() for p in pieces(raw)):
            m = _HEADING.match(line.strip())
            if m:
                flush()
                level = len(m.group(1))
                path[:] = path[: level - 1]
                path.append(m.group(2).strip()[:120])
                continue
            if not line.strip():
                if buf and buf[-1] != "":
                    buf.append("")
                continue
            buf.append(line.rstrip())
            buf_pages.append(page)  # type: ignore[arg-type]
            if sum(len(x) + 1 for x in buf) >= max_chars:
                tail = "\n".join(buf)[-overlap:] if overlap else ""
                flush()
                if tail:
                    buf.append(tail)
                    buf_pages.append(page)  # type: ignore[arg-type]
    flush()
    return passages


def headings(text: str) -> list[tuple[int, str, str]]:
    """[(level, title, body under it)] for every markdown heading - used to propose tree nodes."""
    out: list[tuple[int, str, list[str]]] = []
    for line in text.splitlines():
        m = _HEADING.match(line.strip())
        if m:
            out.append((len(m.group(1)), m.group(2).strip(), []))
        elif out:
            out[-1][2].append(line)
    return [(lvl, title, "\n".join(body).strip()) for lvl, title, body in out]
