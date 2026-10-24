"""Import and export the knowledge tree: JSON (complete, re-importable), Excel/CSV and Markdown.

Imports never write straight into the tree: they produce a draft change set R&D reviews and publishes.
"""

from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from pathlib import Path
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import KnowledgeDomain, KnowledgeNode, KnowledgeNodeDocument, StoredFile, TreeChangeSet, User, VerifiedDocument
from . import storage
from .extract import headings
from .knowledge_tree import add_change

FORMAT = "pathmind-tree"


class TreeFileError(ValueError):
    pass


# ── Reading files into a nested structure ──────────────────────────────────────
# Each node: {"title", "summary", "body", "children": [...]}


def parse(filename: str, data: bytes) -> list[dict]:
    """Return a list of domains: [{"name", "slug"?, ..., "nodes": [...]}]. Unnamed = the target domain."""
    name = filename.lower()
    if name.endswith(".json"):
        try:
            js = json.loads(data.decode("utf-8-sig"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise TreeFileError("The JSON file is not valid.") from exc
        if isinstance(js, dict) and js.get("format") == FORMAT:
            return [_clean_domain(d) for d in js.get("domains") or []]
        if isinstance(js, dict) and "nodes" in js:
            return [_clean_domain(js)]
        if isinstance(js, list):
            return [{"name": "", "nodes": [_clean_node(n) for n in js]}]
        raise TreeFileError("Unrecognised JSON layout. Export a tree from PathMind to see the format.")
    if name.endswith((".md", ".markdown", ".txt")):
        return [{"name": "", "nodes": _from_markdown(data.decode("utf-8-sig", errors="replace"))}]
    if name.endswith((".xlsx", ".csv")):
        return _from_rows(_rows(name, data))
    raise TreeFileError("Use a .json, .xlsx, .csv or .md file.")


def _clean_node(n: Any) -> dict:
    if not isinstance(n, dict) or not str(n.get("title") or "").strip():
        raise TreeFileError("Every node needs a title.")
    return {"title": str(n["title"]).strip()[:200], "summary": str(n.get("summary") or ""),
            "body": str(n.get("body") or ""), "children": [_clean_node(c) for c in n.get("children") or []]}


def _clean_domain(d: Any) -> dict:
    if not isinstance(d, dict):
        raise TreeFileError("Each domain must be an object.")
    return {"name": str(d.get("name") or "").strip()[:120], "slug": str(d.get("slug") or "").strip()[:64],
            "description": str(d.get("description") or ""), "answer_policy": d.get("answer_policy"),
            "citation_detail": d.get("citation_detail"), "nodes": [_clean_node(n) for n in d.get("nodes") or []]}


def _from_markdown(text: str) -> list[dict]:
    roots: list[dict] = []
    stack: list[tuple[int, dict]] = []
    for level, title, body in headings(text):
        node = {"title": title[:200], "summary": "", "body": body, "children": []}
        while stack and stack[-1][0] >= level:
            stack.pop()
        (stack[-1][1]["children"] if stack else roots).append(node)
        stack.append((level, node))
    if not roots:
        raise TreeFileError("No headings found. Use # Heading lines to mark the topics.")
    return roots


def _rows(name: str, data: bytes) -> list[dict[str, str]]:
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        raw = [["" if v is None else str(v).strip() for v in r] for r in ws.iter_rows(values_only=True)]
    else:
        text = data.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        raw = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), dialect)]
    raw = [r for r in raw if any(r)]
    if not raw:
        raise TreeFileError("The file is empty.")
    head = [h.lower() for h in raw[0]]
    if "path" not in head:
        raise TreeFileError("Columns needed: domain (optional), path ('A > B > C'), summary, body.")
    return [dict(zip(head, r)) for r in raw[1:]]


def _from_rows(rows: list[dict[str, str]]) -> list[dict]:
    domains: dict[str, dict] = {}
    for r in rows:
        parts = [p.strip() for p in re.split(r"\s*>\s*", r.get("path", "")) if p.strip()]
        if not parts:
            continue
        dom = domains.setdefault(r.get("domain", ""), {"name": r.get("domain", ""), "nodes": []})
        level = dom["nodes"]
        node = None
        for part in parts:
            node = next((n for n in level if n["title"] == part[:200]), None)
            if node is None:
                node = {"title": part[:200], "summary": "", "body": "", "children": []}
                level.append(node)
            level = node["children"]
        if node is not None:
            node["summary"] = r.get("summary", "") or node["summary"]
            node["body"] = r.get("body", "") or node["body"]
    if not domains:
        raise TreeFileError("No rows with a path.")
    return list(domains.values())


def to_change_set(db: Session, domain: KnowledgeDomain, nodes: list[dict], actor: User, title: str,
                  bulk_import_id: str | None = None, parent_id: str | None = None) -> TreeChangeSet:
    cs = TreeChangeSet(domain_id=domain.id, title=title[:200], origin="import", created_by=actor.id,
                       bulk_import_id=bulk_import_id)
    db.add(cs)
    db.flush()
    counter = [0]

    def walk(items: list[dict], parent_key: str | None):
        for pos, n in enumerate(items, start=1):
            counter[0] += 1
            key = f"i{counter[0]}"
            payload = {"title": n["title"], "summary": n.get("summary", ""), "body": n.get("body", "")}
            if parent_key:
                payload["position"] = pos  # top-level items go after what the parent already has
                payload["parent_temp_key"] = parent_key
            elif parent_id:
                payload["parent_id"] = parent_id
            add_change(db, cs, "add_node", temp_key=key, payload=payload, origin="suggested")
            walk(n.get("children") or [], key)

    walk(nodes, None)
    return cs


# ── Export ────────────────────────────────────────────────────────────────────


def _nested(db: Session, domain: KnowledgeDomain, with_docs: bool) -> list[dict]:
    nodes = db.scalars(select(KnowledgeNode).where(KnowledgeNode.domain_id == domain.id, KnowledgeNode.status == "published")
                       .order_by(KnowledgeNode.position, KnowledgeNode.title)).all()
    by_parent: dict[str | None, list[KnowledgeNode]] = {}
    for n in nodes:
        by_parent.setdefault(n.parent_id, []).append(n)

    def build(pid):
        out = []
        for n in by_parent.get(pid, []):
            item: dict[str, Any] = {"title": n.title, "summary": n.summary, "body": n.body, "children": build(n.id)}
            if with_docs:
                docs = db.scalars(select(VerifiedDocument).join(KnowledgeNodeDocument, KnowledgeNodeDocument.document_id == VerifiedDocument.id)
                                  .where(KnowledgeNodeDocument.node_id == n.id, VerifiedDocument.status == "verified")).all()
                item["documents"] = [{"id": d.id, "title": d.title, "doc_version": d.doc_version} for d in docs]
            out.append(item)
        return out

    return build(None)


def export_json(db: Session, domains: list[KnowledgeDomain]) -> bytes:
    data = {"format": FORMAT, "version": 1, "domains": [
        {"slug": d.slug, "name": d.name, "description": d.description, "answer_policy": d.answer_policy,
         "citation_detail": d.citation_detail, "nodes": _nested(db, d, True)} for d in domains]}
    return json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8")


def _flat(nodes: list[dict], prefix: list[str]) -> list[tuple[str, dict]]:
    out = []
    for n in nodes:
        path = prefix + [n["title"]]
        out.append((" > ".join(path), n))
        out += _flat(n["children"], path)
    return out


def export_xlsx(db: Session, domains: list[KnowledgeDomain]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Tree"
    ws.append(["domain", "path", "summary", "body"])
    for d in domains:
        for path, n in _flat(_nested(db, d, False), []):
            ws.append([d.name, path, n["summary"], n["body"]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_markdown(db: Session, domains: list[KnowledgeDomain]) -> bytes:
    lines: list[str] = []

    def walk(nodes, level):
        for n in nodes:
            lines.append("#" * min(level, 6) + " " + n["title"])
            if n["summary"]:
                lines.append(f"> {n['summary']}")
            if n["body"]:
                lines.append(n["body"])
            lines.append("")
            walk(n["children"], level + 1)

    for d in domains:
        lines.append(f"# {d.name}")
        if d.description:
            lines.append(d.description)
        lines.append("")
        walk(_nested(db, d, False), 2)
    return "\n".join(lines).encode("utf-8")


def export_zip(db: Session, domains: list[KnowledgeDomain], data_dir: Path) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("tree.json", export_json(db, domains))
        z.writestr("tree.md", export_markdown(db, domains))
        for d in domains:
            docs = db.scalars(select(VerifiedDocument).where(VerifiedDocument.domain_id == d.id, VerifiedDocument.status == "verified"))
            for doc in docs:
                if doc.file_id:
                    f = db.get(StoredFile, doc.file_id)
                    if f is not None and f.deleted_at is None:
                        try:
                            z.writestr(f"documents/{d.slug}/{doc.id[:8]}-{storage.safe_name(f.filename)}",
                                       storage.read_bytes(db, data_dir, f))
                            continue
                        except Exception:  # noqa: BLE001 - fall back to the extracted text
                            pass
                z.writestr(f"documents/{d.slug}/{doc.id[:8]}-{storage.safe_name(doc.title)}.txt", doc.text or "")
    return buf.getvalue()
